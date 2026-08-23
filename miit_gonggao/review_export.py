"""公告参数页 PDF -> 公告参数 Excel。

把 `gonggao query --download` 下载的《汽车产品技术参数》PDF（单页、版式固定）解析成
结构化字段，按「公告参数评审模板」的关键参数格式导出：行=参数项、列=各配置型号，
附「备注」列标明参数真实来源。公告页没有的目录参数（通用名称/续驶里程/油耗/电池质量/
储电量）尽力从减免税目录库（data/jianmian_catalog.sqlite）按公告型号补齐。

用法：
    python3 main.py review <车型>                       # 解析 downloads/<车型>/*.pdf
    python3 main.py review downloads/<车型>/xx.pdf -o output/评审.xlsx
"""

from __future__ import annotations

import argparse
import re
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from miit_gonggao.core import save_workbook_atomic

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DOWNLOAD_DIR = PROJECT_ROOT / "downloads"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"
DEFAULT_DB_PATH = PROJECT_ROOT / "data" / "jianmian_catalog.sqlite"

# ---------------------------------------------------------------------------
# PDF 解析：公告参数页为单页固定版式，正文是「标签: 值」流 + 底部多栏交错区。
# 解析核心只依赖词坐标（text/x0/x1/top），便于离线测试。
# ---------------------------------------------------------------------------

# 行内标签 -> 字段 key。按「标签后必跟冒号」tokenize，同一行可有多个标签，
# 值 = 本标签冒号后到下一个标签前的文本。@ 开头是状态标记（外形/货厢的 长宽高 复用）。
LINE_LABELS: list[tuple[str, str]] = [
    ("产品型号名称", "model_full"),
    ("企业名称", "company"),
    ("产品商标", "trademark"),
    ("生产地址", "production_address"),
    ("注册地址", "register_address"),
    ("目录序号", "catalog_seq"),
    ("产品号", "product_no"),
    ("产品ID", "product_id"),
    ("发布日期", "publish_date"),
    ("生效日期", "effective_date"),
    ("外形尺寸(长×宽×高)(mm)", "@dims_outline"),
    ("货厢栏板内尺寸(长×宽×高)(mm)", "@dims_cargo"),
    ("总质量(kg)", "gross_mass"),
    ("轴距(mm)", "wheelbase"),
    ("整备质量(kg)", "curb_mass"),
    ("轴数", "axle_count"),
    ("额定载质量(kg)", "rated_load"),
    ("转向型式", "steering"),
    ("准拖挂车总质量(kg)", "trailer_mass"),
    ("半挂车鞍座最大允许承载质量(kg)", "saddle_mass"),
    ("防抱死系统", "abs"),
    ("载质量利用系数", "load_ratio"),
    ("最高车速(km/h)", "top_speed"),
    ("额定载客(含驾驶员)(人)", "passengers"),
    ("驾驶室准乘人数(人)", "cab_passengers"),
    ("前悬/后悬(mm)", "overhang"),
    ("接近角/离去角(。)", "angles"),
    ("接近角/离去角(°)", "angles"),
    ("前轮距(mm)", "front_track"),
    ("钢板弹簧片数(前/后)", "springs"),
    ("后轮距(mm)", "rear_track"),
    ("轮胎数", "tire_count"),
    ("轮胎规格", "tire_spec"),
    ("燃料种类", "fuel_type"),
    ("排放依据标准", "emission_standard"),
    ("轴荷", "axle_load"),
    ("长", "@dim_length"),
    ("宽", "@dim_width"),
    ("高", "@dim_height"),
]

_LABEL_PATTERN = re.compile(
    "(?P<label>"
    + "|".join(re.escape(label) for label in sorted({label for label, _ in LINE_LABELS}, key=len, reverse=True))
    + r")\s*[:：]"
)
_LABEL_KEY = dict(LINE_LABELS)

# 发动机块表头列：标签前缀 -> 字段 key（值按 x 坐标归列）
ENGINE_COLUMNS: list[tuple[str, str]] = [
    ("发动机型号", "engine_model"),
    ("发动机生产企业", "engine_maker"),
    ("排量(ml)", "engine_displacement"),
    ("功率(kw)", "engine_power"),
    ("油耗", "fuel_consumption_page"),
    ("车身反光标识说明", "reflective_marking"),
]

# 底部左/中栏分界（页面宽 595pt，实测左栏 <145、中栏 <350、右栏「其他」>=390）
_LEFT_COL_MAX = 145.0
_MID_COL_MAX = 350.0
_RIGHT_COL_MIN = 380.0


def _group_lines(words: list[dict[str, Any]], tolerance: float = 3.0) -> list[list[dict[str, Any]]]:
    lines: list[list[dict[str, Any]]] = []
    for word in sorted(words, key=lambda w: (float(w["top"]), float(w["x0"]))):
        if lines and abs(float(word["top"]) - float(lines[-1][0]["top"])) <= tolerance:
            lines[-1].append(word)
        else:
            lines.append([word])
    for line in lines:
        line.sort(key=lambda w: float(w["x0"]))
    return lines


def _line_text(line: list[dict[str, Any]]) -> str:
    return " ".join(str(w["text"]) for w in line)


def parse_words(words: list[dict[str, Any]]) -> dict[str, str]:
    """从词坐标列表解析公告参数页字段。"""
    fields: dict[str, str] = {}
    lines = _group_lines(words)

    def find_word(prefix: str) -> dict[str, Any] | None:
        for word in words:
            if str(word["text"]).startswith(prefix):
                return word
        return None

    other_label = find_word("其他")
    other_top = float(other_label["top"]) if other_label else float("inf")

    # 1) 行内「标签: 值」tokenize（外形/货厢 长宽高按出现顺序归组）。
    # 「其他」右栏在底部与左栏（轴荷等）同行交错，tokenize 前先剔除右栏词。
    token_lines = _group_lines(
        [w for w in words if not (float(w["x0"]) >= _RIGHT_COL_MIN and float(w["top"]) > other_top - 3)]
    )
    dims_group = "outline"
    for line in token_lines:
        text = _line_text(line)
        matches = list(_LABEL_PATTERN.finditer(text))
        for index, match in enumerate(matches):
            key = _LABEL_KEY[match.group("label")]
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            value = text[match.end() : end].strip(" ，,")
            if key == "@dims_outline":
                dims_group = "outline"
            elif key == "@dims_cargo":
                dims_group = "cargo"
            elif key.startswith("@dim_"):
                fields[f"{dims_group}_{key[5:]}"] = value
            elif value and key not in fields:
                fields[key] = value

    # 2) 批次（无冒号：「批次 403」，粘在生效日期值尾部）
    match = re.search(r"批次[:：\s]*(\d+)", "\n".join(_line_text(line) for line in lines))
    if match:
        fields["batch"] = match.group(1)
    if "effective_date" in fields:
        fields["effective_date"] = re.sub(r"\s*批次.*$", "", fields["effective_date"]).strip()

    # 3) 产品型号名称 -> 型号 + 产品名称
    model_full = fields.get("model_full", "")
    split = re.match(r"^([A-Za-z0-9\-]+)\s*型(.+)$", model_full)
    fields["model_code"] = split.group(1) if split else model_full
    fields["product_name"] = split.group(2).strip() if split else ""

    # 4) 底部多栏交错区：按坐标定位
    # 4a) 发动机块：表头行的各标签 x 坐标定列，下方到「其他」之前的词按最近列归属
    engine_header = find_word("发动机型号")
    if engine_header is not None:
        header_top = float(engine_header["top"])
        header_line = next(
            (line for line in lines if any(w is engine_header for w in line)),
            [],
        )
        column_x: list[tuple[float, str]] = []
        for label, key in ENGINE_COLUMNS:
            for word in header_line:
                if str(word["text"]).startswith(label):
                    column_x.append((float(word["x0"]), key))
                    break
        column_x.sort()
        buckets: dict[str, list[str]] = {key: [] for _, key in column_x}
        for line in lines:
            top = float(line[0]["top"])
            if not (header_top + 2 < top < other_top - 1):
                continue
            for word in line:
                x0 = float(word["x0"])
                owner = None
                for col_x, key in column_x:
                    if x0 >= col_x - 10:
                        owner = key
                if owner:
                    buckets[owner].append(str(word["text"]))
        for key, parts in buckets.items():
            if parts:
                fields[key] = "".join(parts)

    # 4b) 「其他」右栏整段
    if other_label is not None:
        parts = [
            str(w["text"])
            for line in lines
            for w in line
            if float(w["x0"]) >= _RIGHT_COL_MIN and float(w["top"]) > other_top + 1
        ]
        if parts:
            fields["other"] = "".join(parts)

    # 4c) 车辆识别代号（左栏）与 底盘型号（中栏）。
    # 个别 PDF 里 VIN 与「承载式车身」无空隙粘成一个词，按字符类型切分。
    vin_label = find_word("车辆识别代号")
    if vin_label is not None:
        vin_top = float(vin_label["top"])
        vin_parts: list[str] = []
        chassis_tail: list[str] = []
        for line in lines:
            for word in line:
                if float(word["x0"]) < _LEFT_COL_MAX and float(word["top"]) > vin_top + 2:
                    match = re.match(r"^([A-Za-z0-9×\-]*)(.*)$", str(word["text"]))
                    if match and match.group(1):
                        vin_parts.append(match.group(1))
                    if match and match.group(2):
                        chassis_tail.append(match.group(2))
        if vin_parts:
            fields["vin"] = "".join(vin_parts)
        if chassis_tail and not fields.get("chassis"):
            fields["chassis"] = "".join(chassis_tail)
    chassis_label = find_word("底盘型号")
    if chassis_label is not None:
        chassis_top = float(chassis_label["top"])
        parts = [
            str(w["text"])
            for line in lines
            for w in line
            if _LEFT_COL_MAX <= float(w["x0"]) < _MID_COL_MAX and float(w["top"]) > chassis_top + 2
        ]
        if parts:
            fields["chassis"] = "".join(parts)

    # 5) 货厢栏板合并
    cargo = [fields.get(f"cargo_{name}", "") for name in ("length", "width", "height")]
    fields["cargo_dims"] = "×".join(cargo) if any(cargo) else ""
    return fields


def parse_gonggao_pdf(pdf_path: Path | str) -> dict[str, str]:
    """解析单个公告参数页 PDF。"""
    import pdfplumber

    pdf_path = Path(pdf_path)
    with pdfplumber.open(pdf_path) as pdf:
        words: list[dict[str, Any]] = []
        for page in pdf.pages:
            words.extend(page.extract_words(use_text_flow=False, keep_blank_chars=False))
    fields = parse_words(words)
    fields["source_file"] = pdf_path.name
    return fields


# ---------------------------------------------------------------------------
# 减免税目录库补充（公告页不显示的目录参数）
# ---------------------------------------------------------------------------

CATALOG_KEYS = ("common_name", "range_km", "fuel_consumption", "battery_mass", "battery_energy")


def lookup_catalog(db_path: Path, model_code: str) -> dict[str, str]:
    if not model_code or not db_path.exists():
        return {}
    try:
        con = sqlite3.connect(db_path)
        rows = con.execute(
            "SELECT catalog, batch, common_name, range_km, fuel_consumption, battery_mass, battery_energy "
            "FROM catalog_rows WHERE model_code = ?",
            (model_code,),
        ).fetchall()
        con.close()
    except sqlite3.Error:
        return {}
    if not rows:
        return {}

    def sort_key(row: tuple) -> tuple[int, int]:
        catalog = str(row[0] or "")
        try:
            batch = int(str(row[1] or "0"))
        except ValueError:
            batch = 0
        return (1 if "购置税" in catalog else 0, batch)

    best: dict[str, str] = {}
    for row in sorted(rows, key=sort_key, reverse=True):
        for key, value in zip(CATALOG_KEYS, row[2:]):
            if not best.get(key) and value:
                best[key] = str(value)
    return best


# ---------------------------------------------------------------------------
# 公告参数 Excel 导出（行结构对齐 公告参数参数评审模板）
# ---------------------------------------------------------------------------

REMARK_PAGE = "公告页参数"
REMARK_CATALOG = "购置税减免目录相关参数；公告参数（公告页不显示）"

SOURCE_PAGE = "公告参数页"
SOURCE_CATALOG = "减免车辆购置税目录"

# (行名, 字段 key, 备注/真实来源)
TEMPLATE_ROWS: list[tuple[str, str, str]] = [
    ("产品型号", "model_code", REMARK_PAGE),
    ("产品号", "product_no", REMARK_PAGE),
    ("产品ID", "product_id", REMARK_PAGE),
    ("批次", "batch", REMARK_PAGE),
    ("发布日期", "publish_date", REMARK_PAGE),
    ("生效日期", "effective_date", REMARK_PAGE),
    ("企业名称", "company", REMARK_PAGE),
    ("产品商标", "trademark", REMARK_PAGE),
    ("通用名称", "common_name", REMARK_CATALOG),
    ("产品名称", "product_name", REMARK_PAGE),
    ("生产企业地址", "production_address", REMARK_PAGE),
    ("注册地址", "register_address", REMARK_PAGE),
    ("目录序号", "catalog_seq", REMARK_PAGE),
    ("外形尺寸长(mm)", "outline_length", REMARK_PAGE),
    ("外形尺寸宽(mm)", "outline_width", REMARK_PAGE),
    ("外形尺寸高(mm)", "outline_height", REMARK_PAGE),
    ("货厢栏板内尺寸(长×宽×高)(mm)", "cargo_dims", REMARK_PAGE),
    ("总质量(kg)", "gross_mass", REMARK_PAGE),
    ("轴荷", "axle_load", REMARK_PAGE),
    ("整备质量(kg)", "curb_mass", REMARK_PAGE),
    ("额定载质量(kg)", "rated_load", REMARK_PAGE),
    ("载质量利用系数", "load_ratio", REMARK_PAGE),
    ("额定载客(含驾驶员)(座位数)(人)", "passengers", REMARK_PAGE),
    ("驾驶室准乘人数(人)", "cab_passengers", REMARK_PAGE),
    ("接近角/离去角(度)（整备质量）", "angles", REMARK_PAGE),
    ("最高车速(km/h)", "top_speed", REMARK_PAGE),
    ("前悬/后悬(mm)", "overhang", REMARK_PAGE),
    ("底盘型号、类别及生产企业", "chassis", REMARK_PAGE),
    ("钢板弹簧片数(前/后)", "springs", REMARK_PAGE),
    ("轴数", "axle_count", REMARK_PAGE),
    ("轴距(mm)", "wheelbase", REMARK_PAGE),
    ("前轮距(mm)", "front_track", REMARK_PAGE),
    ("后轮距(mm)", "rear_track", REMARK_PAGE),
    ("轮胎数", "tire_count", REMARK_PAGE),
    ("轮胎规格", "tire_spec", REMARK_PAGE),
    ("转向型式", "steering", REMARK_PAGE),
    ("车辆识别代号（VIN）", "vin", REMARK_PAGE),
    ("燃料种类", "fuel_type", REMARK_PAGE),
    ("排放依据标准", "emission_standard", REMARK_PAGE),
    ("发动机生产企业名称", "engine_maker", REMARK_PAGE),
    ("发动机型号", "engine_model", REMARK_PAGE),
    ("发动机排量(ml)", "engine_displacement", REMARK_PAGE),
    ("发动机额定功率（电动机峰值功率）(kW)", "engine_power", REMARK_PAGE),
    ("燃油消耗量(L/100km)", "fuel_consumption_page", REMARK_PAGE),
    ("车身反光标识说明", "reflective_marking", REMARK_PAGE),
    ("续驶里程（km）", "range_km", REMARK_CATALOG),
    ("燃料消耗量（L/100km）", "fuel_consumption", REMARK_CATALOG),
    ("动力蓄电池包组质量(kg)", "battery_mass", REMARK_CATALOG),
    ("储能装置总储电量（kWh）", "battery_energy", REMARK_CATALOG),
    ("防抱死制动系统（有/无）", "abs", REMARK_PAGE),
    ("准拖挂车总质量(kg)", "trailer_mass", REMARK_PAGE),
    ("半挂车鞍座最大允许承载质量(kg)", "saddle_mass", REMARK_PAGE),
    ("其他", "other", REMARK_PAGE),
]

_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_CATALOG_FILL = PatternFill("solid", fgColor="FFF7E6")
_THIN_SIDE = Side(style="thin", color="BFBFBF")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)


def _source_note(remark: str, values: list[str]) -> str:
    has_value = any(value and value != "N/A" for value in values)
    if remark == REMARK_CATALOG:
        return SOURCE_CATALOG if has_value else f"{SOURCE_CATALOG}未命中"
    if has_value:
        return SOURCE_PAGE
    return f"{SOURCE_PAGE}未显示/未解析"


def export_review_excel(
    parsed_list: list[dict[str, str]],
    output_path: Path,
    *,
    title: str | None = None,
) -> Path:
    """把若干配置的解析结果导出为公告参数 Excel（列=配置）。"""
    if not parsed_list:
        raise ValueError("没有可导出的公告解析结果")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "关键参数"

    total_columns = 1 + len(parsed_list) + 1
    display_name = title or next(
        (p.get("common_name") for p in parsed_list if p.get("common_name")),
        parsed_list[0].get("trademark", ""),
    )
    title_cell = worksheet.cell(row=1, column=1, value=f"{display_name} 公告关键参数评审表".strip())
    title_cell.font = Font(bold=True, size=14)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)

    headers = ["配置\n(公告时此行信息不显示)"] + [
        parsed.get("model_code") or parsed.get("source_file", f"配置{index + 1}")
        for index, parsed in enumerate(parsed_list)
    ] + ["备注"]
    worksheet.append(headers)
    for cell in worksheet[2]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER

    for label, key, remark in TEMPLATE_ROWS:
        values = [parsed.get(key, "") or "N/A" for parsed in parsed_list]
        worksheet.append([label, *values, _source_note(remark, values)])
        current = worksheet.max_row
        for column in range(1, total_columns + 1):
            cell = worksheet.cell(row=current, column=column)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = _THIN_BORDER
            if remark == REMARK_CATALOG and 2 <= column <= 1 + len(parsed_list):
                cell.fill = _CATALOG_FILL
        worksheet.cell(row=current, column=1).font = Font(bold=True)

    worksheet.freeze_panes = "B3"
    worksheet.column_dimensions["A"].width = 34
    for index in range(len(parsed_list)):
        worksheet.column_dimensions[get_column_letter(2 + index)].width = 26
    worksheet.column_dimensions[get_column_letter(total_columns)].width = 34

    source_sheet = workbook.create_sheet("数据来源")
    source_sheet.append(["产品型号", "产品ID", "批次", "PDF 文件", "目录参数"])
    for cell in source_sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(bold=True)
    for parsed in parsed_list:
        source_sheet.append(
            [
                parsed.get("model_code", ""),
                parsed.get("product_id", ""),
                parsed.get("batch", ""),
                parsed.get("source_file", ""),
                "已从减免税目录补充" if parsed.get("common_name") or parsed.get("range_km") else "目录未命中",
            ]
        )
    for column, width in zip("ABCDE", (24, 14, 10, 44, 20)):
        source_sheet.column_dimensions[column].width = width

    save_workbook_atomic(workbook, output_path)
    return output_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _collect_pdfs(inputs: list[str], download_dir: Path) -> tuple[list[Path], str]:
    pdfs: list[Path] = []
    label = ""
    for item in inputs:
        path = Path(item).expanduser()
        if path.is_file() and path.suffix.lower() == ".pdf":
            pdfs.append(path)
        elif path.is_dir():
            pdfs.extend(sorted(path.glob("*.pdf")))
            label = label or path.name
        elif (download_dir / item).is_dir():
            pdfs.extend(sorted((download_dir / item).glob("*.pdf")))
            label = label or item
        else:
            raise SystemExit(f"找不到 PDF 或目录: {item}（也不在 {download_dir} 下）")
    unique: list[Path] = []
    for pdf in pdfs:
        if pdf not in unique:
            unique.append(pdf)
    return unique, label


def _safe_filename_part(value: str) -> str:
    text = str(value or "").strip()
    for char in '<>:"/\\|?*':
        text = text.replace(char, "_")
    return "_".join(text.split()) or "未命名"


def _batch_filename_part(parsed_list: list[dict[str, str]]) -> str:
    """多批次输出区间式标签：单批次「第406批」，多批次「第394-406批」（取最小-最大）。"""
    batches: list[str] = []
    for parsed in parsed_list:
        batch = str(parsed.get("batch") or "").strip()
        if batch and batch not in batches:
            batches.append(batch)
    if not batches:
        return "未知批次"
    numeric = sorted({int(batch) for batch in batches if batch.isdigit()})
    others = [batch for batch in batches if not batch.isdigit()]
    labels: list[str] = []
    if len(numeric) == 1:
        labels.append(f"第{numeric[0]}批")
    elif numeric:
        labels.append(f"第{numeric[0]}-{numeric[-1]}批")
    labels.extend(others)
    return "-".join(labels)


def _next_available_path(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 1000):
        candidate = path.with_name(f"{path.stem}_{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise SystemExit(f"输出文件已存在且无法自动命名: {path}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="main.py review",
        description="把公告参数页 PDF 转换为公告参数 Excel（格式对齐模板）",
    )
    parser.add_argument(
        "inputs",
        nargs="+",
        help="车型名（对应 downloads/<车型>/ 目录）、PDF 文件或目录，可多个",
    )
    parser.add_argument("-o", "--output", help="输出 xlsx 路径，默认 output/公告参数_<车型>_<批次>.xlsx")
    parser.add_argument("--title", help="报表标题中的车型名，默认取目录通用名称")
    parser.add_argument("--db", help=f"减免税目录库路径，默认 {DEFAULT_DB_PATH}")
    parser.add_argument("--no-catalog", action="store_true", help="不从减免税目录库补充参数")
    parser.add_argument("--download-dir", help=f"公告 PDF 根目录，默认 {DEFAULT_DOWNLOAD_DIR}")
    args = parser.parse_args(argv)

    download_dir = Path(args.download_dir).expanduser() if args.download_dir else DEFAULT_DOWNLOAD_DIR
    pdfs, label = _collect_pdfs(args.inputs, download_dir)
    if not pdfs:
        raise SystemExit("未找到任何 PDF。请先运行 fetch/gonggao query --download 下载公告参数页。")

    db_path = Path(args.db).expanduser() if args.db else DEFAULT_DB_PATH
    parsed_list: list[dict[str, str]] = []
    for pdf in pdfs:
        print(f"解析 {pdf.name} ...")
        fields = parse_gonggao_pdf(pdf)
        if not args.no_catalog:
            fields.update({k: v for k, v in lookup_catalog(db_path, fields.get("model_code", "")).items() if v})
        parsed_list.append(fields)

    if args.output:
        output_path = Path(args.output).expanduser()
    else:
        stem = args.title or label or parsed_list[0].get("common_name") or parsed_list[0].get("model_code", "公告")
        output_path = DEFAULT_OUTPUT_DIR / (
            f"公告参数_{_safe_filename_part(stem)}_{_safe_filename_part(_batch_filename_part(parsed_list))}.xlsx"
        )
        output_path = _next_available_path(output_path)

    export_review_excel(parsed_list, output_path, title=args.title or label or None)
    print(f"已导出公告参数表（{len(parsed_list)} 个配置）: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
