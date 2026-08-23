"""Excel 导出。"""

from __future__ import annotations

import re
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
DIFF_FILL = PatternFill("solid", fgColor="FFF2CC")
ADVANTAGE_FILL = PatternFill("solid", fgColor="C6EFCE")
SECTION_FILL = PatternFill("solid", fgColor="EDEDED")
_THIN_SIDE = Side(style="thin", color="C9CDD4")
THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# 值为这些时视为「无此配置」
EMPTY_VALUES = {"", "-", "–", "—", "无", "N/A", "n/a", "null", "None"}


def export_to_excel(
    *,
    metadata: dict[str, str | int],
    model_rows: list[dict[str, str]],
    scalar_rows: list[dict[str, object]],
    color_rows: list[dict[str, str]],
    output_path: Path,
) -> Path:
    """导出完整车型对比 Excel。

    sheet 结构：说明 / 车型信息 / 配置分析 / 详细配置表 / 颜色明细。
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    intro_sheet = workbook.active
    intro_sheet.title = "说明"
    models_sheet = workbook.create_sheet("车型信息")
    analysis_sheet = workbook.create_sheet("配置分析")
    params_sheet = workbook.create_sheet("详细配置表")
    colors_sheet = workbook.create_sheet("颜色明细")

    _write_intro_sheet(intro_sheet, metadata)
    _write_rows(models_sheet, model_rows)
    _write_analysis_sheet(analysis_sheet, model_rows, scalar_rows)
    _write_params_sheet(params_sheet, model_rows, scalar_rows)
    _write_rows(colors_sheet, color_rows)

    for worksheet in (models_sheet, colors_sheet):
        _format_worksheet(worksheet)

    _save_atomic(workbook, output_path)
    return output_path


def _save_atomic(workbook: Workbook, output_path: Path) -> None:
    """先写临时文件再原子替换，避免进程中途退出留下半截 Excel。

    autohome_cc 不依赖 miit_gonggao，故与 miit_gonggao.core.save_workbook_atomic 各存一份。
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        dir=output_path.parent,
        prefix=f".{output_path.stem}.",
        suffix=output_path.suffix,
        delete=False,
    ) as tmp_file:
        tmp_path = Path(tmp_file.name)
    try:
        workbook.save(tmp_path)
        tmp_path.replace(output_path)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def _write_intro_sheet(worksheet, metadata: dict[str, str | int]) -> None:
    worksheet.append(["项目", "内容"])
    worksheet.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for key, value in metadata.items():
        worksheet.append([key, value])
    worksheet.append(
        [
            "sheet 说明",
            "配置分析=按动力总成分组（能源×驱动×电池电量×座位数）的整合参数透视；详细配置表=全量车型参数对比",
        ]
    )
    _style_header_row(worksheet)
    worksheet.freeze_panes = "A2"
    _apply_borders(worksheet)
    _autosize_columns(worksheet, {1: 18, 2: 120})


# ---------------------------------------------------------------------------
# 配置分析
# ---------------------------------------------------------------------------


def _normalize_value(value: object) -> str:
    text = str(value).strip() if value is not None else ""
    return "" if text in EMPTY_VALUES else text


def _is_standard(value: object) -> bool:
    return "●" in str(value)


def _is_optional(value: object) -> bool:
    return "○" in str(value)


# ---------------------------------------------------------------------------
# 动力总成分组：按 能源（燃油/混动/插混/增程/纯电）× 驱动（两驱/四驱）× 电池电量 × 座位数 归组
# ---------------------------------------------------------------------------

# 能源类型取值 -> 简称（按顺序匹配关键词）
_ENERGY_LABELS: list[tuple[str, tuple[str, ...]]] = [
    ("纯电", ("纯电",)),
    ("增程", ("增程",)),
    ("插混", ("插电",)),
    ("混动", ("油电", "混合", "混动")),
    ("燃油", ("汽油", "柴油", "燃油")),
]


def _find_param_values(scalar_rows: list[dict[str, object]], keywords: tuple[str, ...], count: int) -> list[str]:
    """按参数名关键词找首个匹配行，返回各车型取值（找不到返回空串列表）。"""
    for row in scalar_rows:
        param = str(row["param"])
        if any(keyword in param for keyword in keywords):
            values = [_normalize_value(v) for v in row["values"]]
            return values + [""] * (count - len(values))
    return [""] * count


def _powertrain_labels(scalar_rows: list[dict[str, object]], count: int) -> list[str]:
    """派生每个车型的动力总成分组标签，如「增程四驱60度电池6座」「纯电两驱100度电池7座」。"""
    energies = _find_param_values(scalar_rows, ("能源类型", "燃料形式"), count)
    drives = _find_param_values(scalar_rows, ("驱动方式",), count)
    batteries = _find_param_values(scalar_rows, ("电池能量", "电池容量"), count)
    seat_counts = _find_param_values(scalar_rows, ("座位数",), count)

    labels: list[str] = []
    for energy_raw, drive_raw, battery_raw, seats_raw in zip(energies, drives, batteries, seat_counts):
        energy = next((label for label, keys in _ENERGY_LABELS if any(k in energy_raw for k in keys)), "")
        drive = "四驱" if "四驱" in drive_raw else ("两驱" if drive_raw else "")
        battery = ""
        match = re.search(r"\d+(?:\.\d+)?", battery_raw)
        if match:
            battery = f"{float(match.group()):g}度电池"
        seats = ""
        match = re.search(r"\d+", seats_raw)
        if match:
            seats = f"{match.group()}座"
        labels.append(f"{energy}{drive}{battery}{seats}" or "未分组")
    return labels


_NUMERIC_RE = re.compile(r"^-?\d+(?:\.\d+)?$")


def _summarize_group_values(values: list[str]) -> str:
    """把同一分组内多个车型的取值合并成一格说明。

    规则：组内一致 → 原值；配置项（含●/○）不一致 → 「区分高低配」；
    纯数值多值 → 按数值升序 `2995/3200`；其他文本多值 → `/` 连接（过长截断）。
    """
    normalized = [_normalize_value(v) for v in values]
    distinct = list(dict.fromkeys(v if v else "-" for v in normalized))
    if len(distinct) == 1:
        return distinct[0]
    if any("●" in v or "○" in v for v in distinct):
        return "区分高低配"
    numbers = [v for v in distinct if v != "-" and _NUMERIC_RE.match(v)]
    if len(numbers) == len([v for v in distinct if v != "-"]):
        joined = "/".join(f"{float(v):g}" for v in sorted(numbers, key=float))
        return joined if "-" not in distinct else f"{joined}（部分无）"
    text = "/".join(distinct)
    return text if len(text) <= 60 else "/".join(distinct[:3]) + f"…等{len(distinct)}种"


def _write_analysis_sheet(worksheet, model_rows: list[dict[str, str]], scalar_rows: list[dict[str, object]]) -> None:
    model_count = len(model_rows)

    # 动力总成分组，并按分组把车型列重排（组序=首次出现，组内保持原顺序）
    raw_labels = _powertrain_labels(scalar_rows, model_count)
    group_order = list(dict.fromkeys(raw_labels))
    order = [i for group in group_order for i, label in enumerate(raw_labels) if label == group]
    model_rows = [model_rows[i] for i in order]
    labels = [raw_labels[i] for i in order]
    scalar_rows = [
        {**row, "values": [list(row["values"])[i] if i < len(list(row["values"])) else "" for i in order]}
        for row in scalar_rows
    ]

    model_names = [str(row.get("车型", "")) for row in model_rows]
    # 每车独有配置：本车有实质值、其余车均为空/无
    unique_rows: list[tuple[int, dict[str, object]]] = []
    if model_count > 1:
        for row in scalar_rows:
            values = list(row["values"])
            present = [i for i, v in enumerate(values) if _normalize_value(v)]
            if len(present) == 1:
                unique_rows.append((present[0], row))
    unique_counts = [sum(1 for idx, _ in unique_rows if idx == i) for i in range(model_count)]

    def append_block_title(title: str) -> None:
        row_idx = worksheet.max_row + (2 if worksheet.max_row > 1 else 0)
        cell = worksheet.cell(row=row_idx, column=1, value=title)
        cell.font = Font(bold=True, size=12)
        cell.fill = SECTION_FILL

    def append_header(headers: list[str]) -> int:
        worksheet.append(headers)
        _style_header_row(worksheet, worksheet.max_row)
        return worksheet.max_row

    # 块一：车型配置强度（动力总成分组=能源×驱动×电池电量×座位数）
    append_block_title("一、车型配置强度")
    append_header(["动力总成分组", "车型", "厂商指导价", "标配项(●)", "选配项(○)", "无此配置(-)", "独有配置数"])
    for index, model in enumerate(model_rows):
        values = [row["values"][index] if index < len(row["values"]) else "" for row in scalar_rows]
        worksheet.append(
            [
                labels[index],
                model.get("车型", ""),
                model.get("厂商指导价", ""),
                sum(1 for v in values if _is_standard(v)),
                sum(1 for v in values if _is_optional(v)),
                sum(1 for v in values if not _normalize_value(v)),
                unique_counts[index] if index < len(unique_counts) else 0,
            ]
        )

    # 块二：整合参数明细（每个动力总成分组一列，组内多值在单元格内合并总结）
    append_block_title("二、整合参数明细（每分组一列，组内多值已合并）")
    value_start_col = 3  # 分组/参数 之后
    group_members = {group: [i for i, label in enumerate(labels) if label == group] for group in group_order}
    if not scalar_rows:
        worksheet.append(["无参数数据。"])
    else:
        header_row = append_header(["分组", "参数", *group_order])
        for row in scalar_rows:
            values = list(row["values"])
            summaries = [
                _summarize_group_values([values[i] for i in group_members[group]]) for group in group_order
            ]
            worksheet.append([row["section"], row["param"], *summaries])
            current = worksheet.max_row
            for offset, summary in enumerate(summaries):
                cell = worksheet.cell(row=current, column=value_start_col + offset)
                others = summaries[:offset] + summaries[offset + 1 :]
                if len(summaries) > 1 and _is_standard(summary) and all(
                    not _is_standard(other) and not _is_optional(other) for other in others
                ):
                    cell.fill = ADVANTAGE_FILL
                elif len(summaries) > 1 and summary != "-" and any(summary != other for other in others):
                    cell.fill = DIFF_FILL
        worksheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(value_start_col - 1 + len(group_order))}{worksheet.max_row}"
        )

    # 块三：车型独有配置
    if model_count > 1:
        append_block_title("三、车型独有配置（其余车型均无此项）")
        if not unique_rows:
            worksheet.append(["无独有配置项。"])
        else:
            append_header(["车型", "动力总成分组", "分组", "参数", "配置值"])
            for index, row in unique_rows:
                name = model_names[index] if index < len(model_names) else f"车型{index + 1}"
                worksheet.append([name, labels[index], row["section"], row["param"], str(row["values"][index])])

    for row_cells in worksheet.iter_rows(min_row=1):
        for cell in row_cells:
            if cell.font is None or not cell.font.bold:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    _apply_borders(worksheet)
    widths = {1: 22, 2: 28, 3: 12}
    for index in range(4, 4 + max(model_count, 2)):
        widths[index] = 22
    _autosize_columns(worksheet, widths)


def _write_params_sheet(worksheet, model_rows: list[dict[str, str]], scalar_rows: list[dict[str, object]]) -> None:
    """详细配置表：全量参数 × 全部车型。数据区不加底色，仅表头有样式。"""
    headers = ["分组", "参数"] + [row["车型"] for row in model_rows]
    worksheet.append(headers)
    for row in scalar_rows:
        worksheet.append([str(row["section"]), row["param"], *list(row["values"])])
    _style_header_row(worksheet)
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    _apply_borders(worksheet)
    widths = {1: 14, 2: 30}
    for index in range(3, len(headers) + 1):
        widths[index] = 20
    _autosize_columns(worksheet, widths)


def _write_rows(worksheet, rows: list[dict[str, object]]) -> None:
    if not rows:
        worksheet.append(["message"])
        worksheet.append(["no rows"])
        _style_header_row(worksheet)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        _autosize_columns(worksheet)
        return

    headers = list(rows[0].keys())
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
    _style_header_row(worksheet)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    _format_worksheet_body(worksheet)
    _autosize_columns(worksheet)


def _apply_borders(worksheet) -> None:
    """仅给非空单元格加细边框。"""
    for row_cells in worksheet.iter_rows(min_row=1):
        for cell in row_cells:
            if cell.value is not None:
                cell.border = THIN_BORDER


def _format_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    _format_worksheet_body(worksheet)
    _apply_borders(worksheet)
    _autosize_columns(worksheet)


def _format_worksheet_body(worksheet) -> None:
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _style_header_row(worksheet, row: int = 1) -> None:
    font = Font(bold=True)
    for cell in worksheet[row]:
        if cell.value is None:
            continue
        cell.fill = HEADER_FILL
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize_columns(worksheet, widths: dict[int, int] | None = None) -> None:
    custom_widths = widths or {}
    for column_index, width in custom_widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        if column_cells[0].column in custom_widths:
            continue
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        max_length = max((len(value) for value in values), default=0)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)
