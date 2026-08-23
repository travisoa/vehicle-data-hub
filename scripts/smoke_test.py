"""离线 smoke test：验证 Excel 导出链路与减免税目录解析可用。"""

from __future__ import annotations

import tempfile
from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from autohome_cc.exporter.excel_writer import export_to_excel  # noqa: E402


def smoke_jianmian() -> None:
    """对缓存附件跑两条解析路径，拦截行重建/章节标签回归；缓存缺失则跳过。"""
    from miit_gonggao import jianmian

    html_fixture = ROOT_DIR / "downloads/jianmian/12168_公告第404批/购置税目录第28批.html"
    docx_fixture = ROOT_DIR / "downloads/jianmian/11974_公告第402批/购置税目录第26批.word.docx"
    if not html_fixture.exists() or not docx_fixture.exists():
        print(
            "Smoke test (jianmian) skipped: 附件缓存不存在（解析边界已由 tests/ 的最小 fixture 覆盖，"
            "深度校验可先运行 jianmian sync）"
        )
        return

    # html 启发式路径：购置税第28批，含乘用车与商用车
    catalog, batch, rows = jianmian.parse_catalog_html(html_fixture)
    assert "减免车辆购置税" in catalog, catalog
    assert batch == "28", batch
    assert len(rows) > 300, len(rows)
    # 结构性校验（不绑定具体车型）：乘用车表须解析出通用名称，且每行都有车辆型号
    passenger = [r for r in rows if r.get("category") == "乘用车" and r.get("common_name")]
    assert passenger, "乘用车表未解析出通用名称"
    assert all(r.get("model_code") for r in passenger), "存在缺失车辆型号的行"
    assert {r["energy_type"] for r in passenger} <= {"插电式混合动力汽车", "纯电动汽车", "燃料电池汽车"}
    assert {r["category"] for r in rows} >= {"乘用车", "客车", "货车", "专用车"}

    # Word docx 精确路径：购置税第26批（带"，2026年第1期"后缀的大批次）
    catalog, batch, rows = jianmian.parse_catalog_docx(docx_fixture)
    assert "减免车辆购置税" in catalog, catalog
    assert batch == "26", batch
    assert len(rows) > 10000, len(rows)
    print("Smoke test (jianmian) passed: html批次28 / docx批次26 解析正常")


def main() -> int:
    smoke_jianmian()
    with tempfile.TemporaryDirectory(prefix="autohome-smoke-") as temp_dir:
        output_path = Path(temp_dir) / "smoke.xlsx"
        export_to_excel(
            metadata={
                "运行模式": "smoke_test",
                "输入车型": "示例车型A, 示例车型B",
                "车型数量": 2,
            },
            model_rows=[
                {
                    "spec_id": "1001",
                    "车型": "示例车型A",
                    "厂商指导价": "20.98万",
                    "经销商报价": "20.58万",
                    "来源": "smoke_test",
                },
                {
                    "spec_id": "1002",
                    "车型": "示例车型B",
                    "厂商指导价": "22.98万",
                    "经销商报价": "22.58万",
                    "来源": "smoke_test",
                },
            ],
            scalar_rows=[
                {
                    "section": "基本参数",
                    "param": "能源类型",
                    "values": ["插电混动", "纯电动"],
                },
                {
                    "section": "车身",
                    "param": "座位数",
                    "values": ["7座", "7座"],
                },
            ],
            color_rows=[
                {
                    "类别": "外观颜色",
                    "spec_id": "1001",
                    "车型": "示例车型A",
                    "内容": "星夜黑",
                }
            ],
            output_path=output_path,
        )

        workbook = load_workbook(output_path)
        expected_sheets = [
            "说明",
            "车型信息",
            "配置分析",
            "详细配置表",
            "颜色明细",
        ]
        assert workbook.sheetnames == expected_sheets, workbook.sheetnames
        assert workbook["车型信息"]["B2"].value == "示例车型A"
        assert workbook["详细配置表"]["C2"].value == "插电混动"

        print(f"Smoke test passed: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
