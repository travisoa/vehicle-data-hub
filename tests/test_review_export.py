"""公告参数页解析 + 评审 Excel 导出测试（离线，基于词坐标 fixture）。"""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from miit_gonggao.review_export import _batch_filename_part, _collect_pdfs, export_review_excel, lookup_catalog, parse_words

FIXTURE = Path(__file__).parent / "fixtures" / "gonggao_pdf_words.json"


def _load_fields() -> dict[str, str]:
    words = json.loads(FIXTURE.read_text(encoding="utf-8"))
    return parse_words(words)


def test_parse_words_core_fields():
    fields = _load_fields()
    assert fields["model_code"] == "ABC6520AP6HEV1"
    assert fields["product_name"] == "插电式混合动力多用途乘用车"
    assert fields["product_no"] == "FBE40656302"
    assert fields["product_id"] == "AB356619"
    assert fields["batch"] == "406"
    assert fields["publish_date"] == "20260509"
    assert fields["effective_date"] == "20260509"  # 不应带「批次 406」尾巴
    assert fields["company"] == "示例汽车工业有限公司"
    assert fields["trademark"] == "示例牌"
    assert fields["catalog_seq"] == "79"
    # 外形/货厢两组长宽高按出现顺序归组
    assert fields["outline_length"] == "5200"
    assert fields["outline_width"] == "1999"
    assert fields["outline_height"] == "1820"
    assert not fields["cargo_dims"]
    assert fields["gross_mass"] == "3560"
    assert fields["curb_mass"] == "2980"
    assert fields["wheelbase"] == "3075"
    assert fields["top_speed"] == "220"
    assert fields["passengers"] == "6"
    assert fields["tire_spec"] == "255/50R20,265/45R21"
    assert fields["fuel_type"] == "汽油/电混合动力"
    assert fields["emission_standard"] == "GB18352.6-2016国Ⅵ"
    # 底部多栏交错区
    assert fields["axle_load"] == "1600/1960"  # 不应混入右栏「其他」文本
    assert fields["vin"] == "LC0DD4C4×××××××××"
    assert fields["chassis"] == "承载式车身"
    assert fields["engine_model"] == "BYD479ZQA"
    assert fields["engine_maker"] == "示例汽车工业有限公司"
    assert fields["engine_displacement"] == "1995"
    assert fields["engine_power"] == "152"
    assert fields["other"].startswith("该产品为新能源车辆")
    assert "峰值功率" in fields["other"]


@pytest.mark.parametrize('id_parts', [
    [('底盘ID', 54, 85, 0)],
    [('底盘', 54, 74, 0), ('ID', 76, 87, 1)],
    [('底', 54, 64, 0), ('盘', 65, 75, 0), ('I', 76, 79, 1), ('D', 80, 86, 1)],
    [('底盘 ID', 54, 87, 0)],
])
def test_converted_layout_keeps_chassis_separate_from_vin_and_full_width_other(id_parts):
    def word(text, x, y):
        return {'text': text, 'x0': x, 'x1': x + 40, 'top': y}

    words = [word('改装车产品技术参数', 180, 20),
             word('产品型号名称:TEST5381/C6型泡沫消防车', 35, 100),
             word('总质量(kg):', 36, 300), word('37675', 117, 300),
             word('车辆识别代号:', 419, 519),
             *[dict(word(text, x0, 527 + offset), x1=x1) for text, x0, x1, offset in id_parts],
             word('底盘型号', 112, 527),
             word('底盘类别', 231, 527), word('底盘名称', 294, 527),
             word('1、', 35, 547), word('1234567', 54, 547),
             word('TEST1380', 112, 547), word('二类', 231, 547),
             word('载货汽车底盘', 295, 547), word('LTEST123×××××××××', 419, 547),
             word('2、', 35, 567), word('7654321', 54, 567),
             word('TEST1381', 112, 567), word('二类', 231, 567),
             word('载货汽车底盘', 295, 567),
             word('油耗:22.8', 36, 627),
             word('车身反光标识说明:', 36, 649), word('企业:示例企业;', 135, 649),
             word('其他:', 38, 669), word('该车用于消防作业;', 53, 689),
             word('运输介质:汽油,不能作为燃料依据。', 53, 701)]
    fields = parse_words(words)
    assert fields['model_code'] == 'TEST5381/C6'
    assert fields['product_name'] == '泡沫消防车'
    assert fields['chassis'] == 'TEST1380;TEST1381'
    assert fields['vin'] == 'LTEST123×××××××××'
    assert [r['product_id'] for r in json.loads(fields['chassis_references'])] == ['1234567', '7654321']
    assert fields['other'] == '该车用于消防作业;运输介质:汽油,不能作为燃料依据。'
    assert fields['fuel_consumption_page'] == '22.8'
    assert fields['reflective_marking'] == '企业:示例企业;'
    assert 'fuel_type' not in fields

    # VIN 原文为空时，底盘 ID 和型号绝不能成为虚假的 VIN。
    fields = parse_words([w for w in words if w['text'] != 'LTEST123×××××××××'])
    assert fields['vin'] == ''


@pytest.mark.parametrize('id_words', [
    [{'text': '底盘ID', 'x0': 54, 'x1': 85, 'top': 527}],
    [{'text': '底盘', 'x0': 54, 'x1': 74, 'top': 527},
     {'text': 'ID', 'x0': 76, 'x1': 87, 'top': 528}],
])
def test_converted_layout_rejects_incomplete_chassis_header(id_words):
    with pytest.raises(ValueError, match='缺少必要表头'):
        parse_words(id_words)


@pytest.mark.parametrize(('id_x', 'id_y'), [(160, 527), (76, 540)])
def test_converted_layout_does_not_join_distant_words(id_x, id_y):
    fields = parse_words([
        {'text': '底盘', 'x0': 54, 'x1': 74, 'top': 527},
        {'text': 'ID', 'x0': id_x, 'x1': id_x + 11, 'top': id_y},
    ])
    assert 'pdf_layout' not in fields


def test_base_layout_parse_can_explicitly_disable_converted_dispatch():
    fields = parse_words(
        [{'text': '底盘ID', 'x0': 54, 'x1': 85, 'top': 527}],
        detect_converted_layout=False,
    )
    assert 'pdf_layout' not in fields


def test_export_review_excel(tmp_path):
    fields = _load_fields()
    fields["source_file"] = "demo.pdf"
    fields["common_name"] = "示例车型N1"
    fields["range_km"] = "170"
    output_path = tmp_path / "review.xlsx"
    export_review_excel([fields, dict(fields, model_code="ABC6520AP6HEV2")], output_path, title="示例车型N1")

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == ["关键参数", "数据来源"]
    worksheet = workbook["关键参数"]
    assert worksheet["A1"].value == "示例车型N1 公告关键参数评审表"
    assert worksheet["B2"].value == "ABC6520AP6HEV1"
    assert worksheet["C2"].value == "ABC6520AP6HEV2"
    rows = {row[0]: row[1:] for row in worksheet.iter_rows(min_row=3, values_only=True)}
    assert rows["产品型号"][0] == "ABC6520AP6HEV1"
    assert rows["通用名称"][0] == "示例车型N1"
    assert rows["续驶里程（km）"][0] == "170"
    assert rows["车辆识别代号（VIN）"][0] == "LC0DD4C4×××××××××"
    # 备注列按真实来源填写，已删除「是否已评审」列
    assert rows["产品号"][-1] == "公告参数页"
    assert rows["通用名称"][-1] == "减免车辆购置税目录"
    # 公告页没有的值填 N/A
    assert rows["货厢栏板内尺寸(长×宽×高)(mm)"][0] == "N/A"
    assert rows["货厢栏板内尺寸(长×宽×高)(mm)"][-1] == "公告参数页未显示/未解析"


def test_batch_filename_part():
    assert _batch_filename_part([{"batch": "406"}, {"batch": "406"}]) == "第406批"
    # 多批次取最小-最大区间，且与 PDF 出现顺序无关
    assert _batch_filename_part([{"batch": "406"}, {"batch": "394"}, {"batch": "405"}]) == "第394-406批"
    assert _batch_filename_part([{}]) == "未知批次"


def test_collect_pdfs_finds_vehicle_under_brand_and_batch(tmp_path):
    pdf = tmp_path / "示例牌" / "示例车型" / "第408批" / "demo.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"%PDF-1.4 test")

    pdfs, label = _collect_pdfs(["示例车型"], tmp_path)

    assert pdfs == [pdf]
    assert label == "示例车型"


def test_lookup_catalog(tmp_path):
    db_path = tmp_path / "catalog.sqlite"
    con = sqlite3.connect(db_path)
    con.execute(
        "CREATE TABLE catalog_rows (catalog TEXT, batch TEXT, model_code TEXT, common_name TEXT, "
        "range_km TEXT, fuel_consumption TEXT, battery_mass TEXT, battery_energy TEXT)"
    )
    con.executemany(
        "INSERT INTO catalog_rows VALUES (?,?,?,?,?,?,?,?)",
        [
            ("享受车船税减免优惠的车型目录", "82", "ABC6520AP6HEV1", "", "170", "6.95", "413", "46.992"),
            ("减免车辆购置税的新能源汽车车型目录", "28", "ABC6520AP6HEV1", "示例车型N1", "170", "6.95", "413", "46.992"),
        ],
    )
    con.commit()
    con.close()

    result = lookup_catalog(db_path, "ABC6520AP6HEV1")
    assert result["common_name"] == "示例车型N1"  # 购置税目录优先
    assert result["range_km"] == "170"
    assert lookup_catalog(db_path, "不存在的型号") == {}
    assert lookup_catalog(tmp_path / "missing.sqlite", "ABC6520AP6HEV1") == {}
