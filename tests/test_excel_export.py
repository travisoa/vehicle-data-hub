"""autohome Excel 导出链路测试（离线）。"""

from __future__ import annotations

from openpyxl import load_workbook

from autohome_cc.exporter.excel_writer import export_to_excel

EXPECTED_SHEETS = ["说明", "车型信息", "配置分析", "详细配置表", "颜色明细"]


def _sheet_values(worksheet) -> list[list[object]]:
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def test_export_to_excel_roundtrip(tmp_path):
    output_path = tmp_path / "export.xlsx"
    export_to_excel(
        metadata={"运行模式": "pytest", "车型数量": 2},
        model_rows=[
            {"spec_id": "1001", "车型": "示例车型A", "厂商指导价": "20.98万", "经销商报价": "-", "来源": "pytest"},
            {"spec_id": "1002", "车型": "示例车型B", "厂商指导价": "22.98万", "经销商报价": "-", "来源": "pytest"},
        ],
        scalar_rows=[
            {"section": "基本参数", "param": "能源类型", "values": ["插电混动", "纯电动"]},
            {"section": "车身", "param": "座位数", "values": ["7座", "7座"]},
            {"section": "舒适配置", "param": "座椅按摩", "values": ["●", "-"]},
            {"section": "智能配置", "param": "激光雷达", "values": ["-", "●"]},
        ],
        color_rows=[{"类别": "外观颜色", "spec_id": "1001", "车型": "示例车型A", "内容": "星夜黑"}],
        output_path=output_path,
    )

    workbook = load_workbook(output_path)
    assert workbook.sheetnames == EXPECTED_SHEETS
    assert workbook["车型信息"]["B2"].value == "示例车型A"
    # 详细配置表：分组/参数后直接是车型列（已无「是否差异」列）
    assert workbook["详细配置表"]["C1"].value == "示例车型A"
    assert workbook["详细配置表"]["C2"].value == "插电混动"

    analysis = _sheet_values(workbook["配置分析"])
    text = "\n".join("|".join(str(c) for c in row if c is not None) for row in analysis)
    # 已删除对比总览/动力总成分组两块，首块即车型配置强度
    assert "对比总览" not in text
    assert "一、车型配置强度" in text
    # 配置强度表带分组标签（能源类型+座位数派生）
    assert any(row[:2] == ["插混7座", "示例车型A"] for row in analysis)
    assert any(row[:2] == ["纯电7座", "示例车型B"] for row in analysis)
    # 整合参数明细保留全参数：差异项和一致项都进入配置分析
    assert "能源类型" in text
    assert any(row[1] == "座位数" and row[0] == "车身" for row in analysis if len(row) > 1 and row[0])
    # 独有配置：A 独有座椅按摩、B 独有激光雷达（第2列为动力总成分组）
    assert any(row[:4] == ["示例车型A", "插混7座", "舒适配置", "座椅按摩"] for row in analysis)
    assert any(row[:4] == ["示例车型B", "纯电7座", "智能配置", "激光雷达"] for row in analysis)


def test_powertrain_grouping(tmp_path):
    """按 能源×驱动×电池电量 归组：标签、每分组一列、组内多值合并总结。"""
    output_path = tmp_path / "group.xlsx"
    export_to_excel(
        metadata={"运行模式": "pytest"},
        model_rows=[
            {"spec_id": "1", "车型": "增程Max四驱", "厂商指导价": "50.98万", "经销商报价": "-", "来源": "t"},
            {"spec_id": "2", "车型": "纯电Max两驱", "厂商指导价": "46.98万", "经销商报价": "-", "来源": "t"},
            {"spec_id": "3", "车型": "增程Ultra四驱", "厂商指导价": "56.98万", "经销商报价": "-", "来源": "t"},
        ],
        scalar_rows=[
            {"section": "基本参数", "param": "能源类型", "values": ["增程式", "纯电动", "增程式"]},
            {"section": "底盘转向", "param": "驱动方式", "values": ["双电机四驱", "前置前驱", "双电机四驱"]},
            {"section": "电池/充电", "param": "电池能量(kWh)", "values": ["56", "100", "56"]},
            {"section": "车身", "param": "座位数(个)", "values": ["6", "7", "6"]},
            # 组间差异：两个增程车一致、与纯电车不同
            {"section": "基本参数", "param": "油箱容积(L)", "values": ["63", "-", "63"]},
            # 组内数值多值：合并为 2995/3200
            {"section": "车身", "param": "总质量(kg)", "values": ["2995", "3060", "3200"]},
            # 组内配置差异：合并为「区分高低配」
            {"section": "舒适配置", "param": "座椅按摩", "values": ["-", "-", "●"]},
        ],
        color_rows=[],
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    analysis = _sheet_values(workbook["配置分析"])
    # 配置强度表按分组排列（同组相邻）
    strength_rows = [row for row in analysis if row[0] in ("增程四驱56度电池6座", "纯电两驱100度电池7座") and row[1] in
                     ("增程Max四驱", "增程Ultra四驱", "纯电Max两驱")]
    assert [row[1] for row in strength_rows] == ["增程Max四驱", "增程Ultra四驱", "纯电Max两驱"]
    # 整合参数明细：每分组一列（增程组在前），组内多值合并
    header = next(row for row in analysis if row[0] == "分组" and row[1] == "参数")
    assert header[2:4] == ["增程四驱56度电池6座", "纯电两驱100度电池7座"]
    fuel_tank = next(row for row in analysis if row[1] == "油箱容积(L)")
    assert fuel_tank[2:4] == ["63", "-"]
    mass = next(row for row in analysis if row[1] == "总质量(kg)")
    assert mass[2:4] == ["2995/3200", "3060"]
    massage = next(row for row in analysis if row[1] == "座椅按摩")
    assert massage[2:4] == ["区分高低配", "-"]
    seats = next(row for row in analysis if row[1] == "座位数(个)")
    assert seats[2:4] == ["6", "7"]


def test_export_single_model_analysis(tmp_path):
    output_path = tmp_path / "single.xlsx"
    export_to_excel(
        metadata={"运行模式": "pytest"},
        model_rows=[{"spec_id": "1001", "车型": "单车型", "厂商指导价": "20万", "经销商报价": "-", "来源": "pytest"}],
        scalar_rows=[{"section": "基本参数", "param": "能源类型", "values": ["纯电动"]}],
        color_rows=[],
        output_path=output_path,
    )
    workbook = load_workbook(output_path)
    assert workbook.sheetnames == EXPECTED_SHEETS
    text = "\n".join(
        "|".join(str(c) for c in row if c is not None) for row in workbook["配置分析"].iter_rows(values_only=True)
    )
    assert "整合参数明细" in text
    assert "能源类型" in text
